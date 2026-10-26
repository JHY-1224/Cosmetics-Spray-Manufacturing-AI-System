import io
import os
import zipfile
from pathlib import Path

import boto3
from botocore.config import Config

from .security import digest, endpoint, fail

TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".csv": "text/csv",
    ".json": "application/json",
    ".wav": "audio/wav",
    ".mp3": "audio/mpeg",
    ".mp4": "video/mp4",
}


def validate_file(name, data, limit):
    ext = Path(name).suffix.lower()
    if not data or len(data) > limit or ext not in TYPES:
        fail(422, "file_invalid", "文件为空、超过大小限制或类型不支持")
    if (
        ext == ".png"
        and not data.startswith(b"\x89PNG\r\n\x1a\n")
        or ext in (".jpg", ".jpeg")
        and not data.startswith(b"\xff\xd8\xff")
        or ext == ".pdf"
        and not data.startswith(b"%PDF-")
        or ext == ".webp"
        and not (data.startswith(b"RIFF") and data[8:12] == b"WEBP")
    ):
        fail(422, "file_magic", "文件内容与类型不符")
    if (ext == ".wav" and not (data.startswith(b"RIFF") and data[8:12] == b"WAVE")
        or ext == ".mp3" and not (data.startswith(b"ID3") or len(data)>2 and data[0]==255 and data[1]&224==224)
        or ext == ".mp4" and not (len(data)>12 and data[4:8]==b"ftyp")):
        fail(422, "file_magic", "音视频原件格式与类型不符；不提供自动识别或训练")
    if ext == ".docx":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if (
                    sum(i.file_size for i in z.infolist()) > 50 * 1024**2
                    or len(z.infolist()) > 2000
                    or "word/document.xml" not in z.namelist()
                ):
                    raise ValueError("压缩包内容或大小非法")
                if any(".." in Path(i.filename).parts or i.filename.startswith("/") for i in z.infolist()):
                    raise ValueError("非法路径")
        except Exception:
            fail(422, "docx_invalid", "DOCX格式错误或解压大小超过限制")
    return TYPES[ext]


class Storage:
    def __init__(self, settings):
        self.settings = settings
        self.root = settings.data_dir / "objects"
        self.root.mkdir(parents=True, exist_ok=True, mode=0o700)
        self.client = None
        if settings.storage_backend == "s3":
            endpoint(settings.s3_endpoint, settings)
            self.client = boto3.client(
                "s3",
                endpoint_url=settings.s3_endpoint,
                aws_access_key_id=settings.s3_access_key,
                aws_secret_access_key=settings.s3_secret_key,
                region_name="us-east-1",
                config=Config(
                    signature_version="s3v4",
                    connect_timeout=5,
                    read_timeout=15,
                    retries={"max_attempts": 2},
                    proxies={},
                    s3={"addressing_style": "path"},
                ),
            )
            from .sdk_control import pin_s3

            pin_s3(self.client, settings)

    def path(self, key):
        p = (self.root / key).resolve()
        if not p.is_relative_to(self.root.resolve()):
            fail(422, "path_invalid", "对象路径越界")
        return p

    def put(self, key, data, media_type):
        used = sum(p.stat().st_size for p in self.root.rglob("*") if p.is_file())
        if used + len(data) > self.settings.cache_max_bytes:
            fail(507, "cache_full", "本地证据缓存已满，拒绝新增且不删除原始资料")
        # Local immutable spool remains available during a center/S3 outage.
        p = self.path(key)
        p.parent.mkdir(parents=True, exist_ok=True)
        with os.fdopen(os.open(p, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600), "wb") as f:
            f.write(data)
            f.flush()
            os.fsync(f.fileno())
        if self.client:
            try:
                self.reconcile(key, media_type, digest(data))
            except ValueError:
                return "conflict"
            except Exception:
                return "spooled"
        return "available"

    def get(self, key, expected_hash):
        p = self.path(key)
        data = (
            p.read_bytes()
            if p.exists()
            else self.client.get_object(Bucket=self.settings.s3_bucket, Key=key)["Body"].read(
                self.settings.max_upload_bytes + 1
            )
        )
        if digest(data) != expected_hash:
            fail(409, "evidence_corrupt", "证据完整性校验失败")
        return data

    def reconcile(self, key, media_type, expected_hash):
        from botocore.exceptions import ClientError

        try:
            existing = self.client.head_object(Bucket=self.settings.s3_bucket, Key=key)
            if existing.get("Metadata", {}).get("sha256") != expected_hash:
                raise ValueError("remote_object_integrity_conflict")
            return
        except ClientError as exc:
            if exc.response["Error"]["Code"] not in ("404", "NoSuchKey", "NotFound"):
                raise
        data = self.get(key, expected_hash)
        self.client.put_object(
            Bucket=self.settings.s3_bucket,
            Key=key,
            Body=data,
            ContentType=media_type,
            Metadata={"sha256": expected_hash},
            IfNoneMatch="*",
        )


def parse_document(name, data):
    ext = Path(name).suffix.lower()
    if ext in (".txt", ".md"):
        return {"status": "parsed", "text": data.decode("utf-8"), "locations": [{"section": "全文"}]}
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if len(reader.pages) > 300:
            fail(422, "pages_limit", "PDF页数超出限制")
        pages = [{"page": i + 1, "text": p.extract_text() or ""} for i, p in enumerate(reader.pages)]
        content = "\n".join(p["text"] for p in pages)
        return {"status": "parsed" if content.strip() else "needs_ocr", "text": content, "locations": pages}
    if ext == ".docx":
        from docx import Document

        doc = Document(io.BytesIO(data))
        paragraphs = [p.text for p in doc.paragraphs]
        paragraphs += [" | ".join(c.text for c in r.cells) for t in doc.tables for r in t.rows]
        content = "\n".join(paragraphs)
        return {
            "status": "parsed" if content.strip() else "needs_review",
            "text": content,
            "locations": [{"paragraph": i + 1, "text": t} for i, t in enumerate(paragraphs)],
        }
    fail(422, "knowledge_type", "知识导入支持TXT/Markdown/PDF/DOCX")
